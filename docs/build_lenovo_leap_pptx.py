"""Lenovo LEAP jury deck for PHISHEYE — dark product UI, real screenshots."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / "docs" / "_leap_media"
OUT = ROOT / "PHISHEYE_Lenovo_LEAP_Jury_Deck.pptx"
DOWNLOADS = Path.home() / "Downloads" / "PHISHEYE_Lenovo_LEAP_Jury_Deck.pptx"

BG = RGBColor(0x08, 0x08, 0x08)
CARD = RGBColor(0x12, 0x12, 0x12)
RED = RGBColor(0xE1, 0x06, 0x00)
WHITE = RGBColor(0xF5, 0xF5, 0xF5)
MUTED = RGBColor(0x9A, 0x9A, 0x9A)
LINE = RGBColor(0x2A, 0x2A, 0x2A)
W, H = Inches(13.333), Inches(7.5)


def set_run(run, text, size=14, bold=False, color=WHITE, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)


def add_text(slide, l, t, w, h, lines, size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        set_run(run, line, size=size, bold=bold, color=color, font=font)
    return box


def add_shape(slide, l, t, w, h, fill=CARD, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def paint(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Inches(0.08), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), W, Inches(0.04))
    top.fill.solid()
    top.fill.fore_color.rgb = RED
    top.line.fill.background()


def footer(slide, page, total=12):
    add_text(slide, 0.45, 7.15, 8, 0.28, ["PHISHEYE  ·  Lenovo LEAP  ·  Scan before you trust"], size=10, color=MUTED)
    add_text(slide, 11.4, 7.15, 1.5, 0.28, [f"{page:02d}  /  {total:02d}"], size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def fit_image(slide, path, l, t, w, h):
    from PIL import Image

    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    box_w, box_h = Inches(w), Inches(h)
    scale = min(box_w / iw, box_h / ih)
    nw, nh = iw * scale, ih * scale
    left = Inches(l) + (box_w - nw) / 2
    top = Inches(t) + (box_h - nh) / 2
    return slide.shapes.add_picture(str(path), left, top, nw, nh)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint(slide)
    return slide


def card(slide, l, t, w, h, title, body, kicker=""):
    add_shape(slide, l, t, w, h, CARD, LINE)
    add_shape(slide, l, t, 0.07, h, RED)
    y = t + 0.16
    if kicker:
        add_text(slide, l + 0.22, y, w - 0.35, 0.22, [kicker.upper()], size=10, color=RED, bold=True)
        y += 0.26
    add_text(slide, l + 0.22, y, w - 0.35, 0.4, [title], size=16, bold=True, color=WHITE)
    add_text(slide, l + 0.22, y + 0.42, w - 0.35, h - (y - t) - 0.55, body, size=13, color=MUTED)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 01 title
    s = new_slide(prs)
    add_text(s, 0.55, 1.35, 12, 0.3, ["LENOVO LEAP  ·  JURY PRESENTATION"], size=13, bold=True, color=RED)
    add_text(s, 0.55, 1.85, 12, 1.2, ["PHISHEYE"], size=60, bold=True, color=WHITE, font="Calibri")
    add_text(s, 0.55, 3.15, 11, 0.7, ["Scan before you trust."], size=28, color=WHITE)
    add_text(
        s,
        0.55,
        3.95,
        11.5,
        1.1,
        [
            "A working phishing and lookalike-domain detector for Indian users.",
            "Dashboard + Chrome extension share one FastAPI engine.",
        ],
        size=18,
        color=MUTED,
    )
    add_shape(s, 0.55, 5.35, 2.1, 0.08, RED)
    add_text(s, 0.55, 5.6, 12, 0.7, ["Live product  ·  React + FastAPI + SQLite  ·  Chrome MV3 scan-before-open gate"], size=14, color=MUTED)
    footer(s, 1)
    notes(s, "Open with the OTP/UPI story, then say the product name once: PHISHEYE. Do not say Flask, Postgres, or CNN.")

    # 02 problem
    s = new_slide(prs)
    add_text(s, 0.55, 0.28, 12, 0.25, ["01  /  THE PROBLEM"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.58, 12, 0.55, ["A padlock does not mean it is SBI."], size=28, bold=True)
    card(s, 0.55, 1.45, 3.9, 4.9, "The SMS arrives", [
        "“Your SBI account will be blocked. Complete KYC in 10 minutes. Share OTP.”",
        "",
        "The page looks real. It has HTTPS. The victim types the OTP. Money leaves.",
    ], "India")
    card(s, 4.7, 1.45, 3.9, 4.9, "Why people fail", [
        "Blacklists are late. A new fake domain can steal for hours.",
        "",
        "Short links and QR codes hide the real host.",
        "",
        "Parents and first-time UPI users need plain English, not certificate jargon.",
    ], "Human")
    card(s, 8.85, 1.45, 3.9, 4.9, "What we must catch", [
        "paypa1.com",
        "sbi-login.xyz",
        "amaz0n-secure.tk",
        "hdfc-login.vercel.app",
        "",
        "Lookalikes that imitate genuine brands — the core of this problem.",
    ], "Lookalike")
    footer(s, 2)
    notes(s, "Speak slowly. Point at sbi-login.xyz. Say HTTPS is cheap. Then go to the product.")

    # 03 idea
    s = new_slide(prs)
    add_text(s, 0.55, 0.28, 12, 0.25, ["02  /  THE IDEA"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.58, 12, 0.5, ["One brain. Two surfaces. Act before the fake page loads."], size=24, bold=True)
    card(s, 0.55, 1.4, 6.05, 2.35, "Dashboard", [
        "Paste a URL, message, email, WhatsApp/SMS text, screenshot, or QR.",
        "Get Simple View (do not enter OTP) and Technical View (evidence).",
    ], "Website")
    card(s, 6.8, 1.4, 5.95, 2.35, "Chrome extension", [
        "Holds the tab. Scores the URL first.",
        "LOW opens. MEDIUM asks. HIGH and CRITICAL stay closed.",
    ], "Main product")
    card(s, 0.55, 3.95, 4.0, 2.45, "Normalize", ["Add https, drop tricks in www and query strings, unwrap short links after an SSRF guard."], "01")
    card(s, 4.7, 3.95, 4.0, 2.45, "Analyze + fuse", ["Feeds, lookalikes, age, TLS, psychology. Confirmed intel cannot be talked down."], "02")
    card(s, 8.85, 3.95, 3.9, 2.45, "Explain + act", ["Plain-English advice. Store in SQLite. Gate the browser."], "03")
    footer(s, 3)
    notes(s, "The extension is the main product. The website is the console. Same POST to FastAPI.")

    # 04 landing
    s = new_slide(prs)
    add_text(s, 0.55, 0.22, 12, 0.22, ["03  /  PRODUCT  ·  LANDING"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.48, 12, 0.35, ["Working UI. Not a mockup."], size=20, bold=True)
    add_shape(s, 0.45, 0.95, 12.45, 6.05, CARD, LINE)
    fit_image(s, MEDIA / "01_landing.png", 0.55, 1.05, 12.25, 5.85)
    footer(s, 4)
    notes(s, "This is the live PHISHEYE landing. Analyze URL, screenshot, Get browser protection.")

    # 05 dashboard + scan
    s = new_slide(prs)
    add_text(s, 0.55, 0.22, 12, 0.22, ["04  /  SECURITY CONSOLE"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.48, 12, 0.32, ["Signed-in operator console. FastAPI + SQLite behind it."], size=16, bold=True)
    add_shape(s, 0.4, 0.95, 6.2, 5.95, CARD, LINE)
    fit_image(s, MEDIA / "02_dashboard.png", 0.5, 1.05, 6.0, 5.75)
    add_shape(s, 6.75, 0.95, 6.1, 5.95, CARD, LINE)
    fit_image(s, MEDIA / "05_scan.png", 6.85, 1.05, 5.9, 5.75)
    footer(s, 5)
    notes(s, "Left: dashboard URL scan. Right: multi-channel scan — URL, message, email, WhatsApp/SMS, screenshot, QR.")

    # 06 channels
    s = new_slide(prs)
    add_text(s, 0.55, 0.28, 12, 0.25, ["05  /  WHAT YOU CAN SCAN"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.58, 12, 0.4, ["Phishing does not arrive as a clean URL. Neither does our input."], size=20, bold=True)
    items = [
        ("URL", "Single address through check-url"),
        ("Message", "Pasted scam text, links extracted"),
        ("Email", "Visible links, no inbox OAuth"),
        ("WhatsApp / SMS", "Paste the text. We do not read the phone."),
        ("Screenshot", "Image of a page or poster"),
        ("QR", "OpenCV reads the code. We do not visit the destination."),
    ]
    for i, (title, body) in enumerate(items):
        col, row = i % 3, i // 3
        card(s, 0.55 + col * 4.2, 1.25 + row * 2.7, 4.0, 2.5, title, [body], f"{i+1:02d}")
    footer(s, 6)
    notes(s, "Be explicit: no Gmail OAuth, no native WhatsApp, QR does not fetch the phishing page.")

    # 07 history + reports
    s = new_slide(prs)
    add_text(s, 0.55, 0.22, 12, 0.22, ["06  /  REAL CHECKS, STORED"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.48, 12, 0.32, ["History is SQLite. Community labels are a signal, not proof."], size=16, bold=True)
    add_shape(s, 0.4, 0.95, 6.35, 5.95, CARD, LINE)
    fit_image(s, MEDIA / "07_history.png", 0.5, 1.05, 6.15, 5.75)
    add_shape(s, 6.9, 0.95, 5.95, 5.95, CARD, LINE)
    fit_image(s, MEDIA / "10_reports.png", 7.0, 1.05, 5.75, 5.75)
    footer(s, 7)
    notes(s, "If they ask about paypal.com SAFE in community reports: users can be wrong. Reports do not override feeds.")

    # 08 analytics
    s = new_slide(prs)
    add_text(s, 0.55, 0.22, 12, 0.22, ["07  /  LIVE CONSOLE METRICS"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.48, 12, 0.32, ["From this prototype’s SQLite — not a published accuracy paper."], size=16, bold=True)
    add_shape(s, 0.4, 0.95, 6.2, 5.95, CARD, LINE)
    fit_image(s, MEDIA / "08_statistics.png", 0.5, 1.05, 6.0, 5.75)
    add_shape(s, 6.75, 0.95, 6.1, 5.95, CARD, LINE)
    fit_image(s, MEDIA / "09_pie.png", 6.85, 1.05, 5.9, 5.75)
    footer(s, 8)
    notes(s, "958 scans / 117 threats / 87.8% safety rate are dashboard numbers from stored checks. Do NOT say 99% or F1.")

    # 09 extension
    s = new_slide(prs)
    add_text(s, 0.55, 0.22, 12, 0.22, ["08  /  THE MAIN PRODUCT  ·  EXTENSION"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.48, 12, 0.32, ["Scan before the other site opens. Same engine as the dashboard."], size=16, bold=True)
    add_shape(s, 0.4, 0.95, 12.55, 5.95, CARD, LINE)
    fit_image(s, MEDIA / "13_extension.png", 0.5, 1.05, 12.35, 5.75)
    footer(s, 9)
    notes(s, "Load unpacked, not Chrome Web Store. Demo: github.com opens if LOW. sbi-login.xyz stays closed. Do not open a live kit.")

    # 10 novelty
    s = new_slide(prs)
    add_text(s, 0.55, 0.28, 12, 0.25, ["09  /  NOVELTY  ·  WHAT IS ACTUALLY NEW"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.58, 12, 0.4, ["Not ‘the first phishing detector’. The combination is the product."], size=20, bold=True)
    points = [
        ("Scan-first gate", "The tab waits. LOW auto-opens. HIGH/CRITICAL never load the kit."),
        ("Intel floors", "A confirmed OpenPhish / GSB hit cannot be vetoed by pretty heuristics. A miss is not SAFE."),
        ("One engine", "Dashboard and extension call the same FastAPI routes. No second fake scorer."),
        ("India language", "Hinglish urgency, UPI/OTP/KYC phrasing, SBI and bank lookalikes."),
        ("Explainable", "Simple View for the user. Technical View for the jury. No black-box 99%."),
        ("Honest privacy", "URL in, advice out. No passwords, no OTP, no SMS inbox, no Gmail OAuth."),
    ]
    for i, (title, body) in enumerate(points):
        col, row = i % 3, i // 3
        card(s, 0.55 + col * 4.2, 1.2 + row * 2.75, 4.0, 2.55, title, [body], f"{i+1:02d}")
    footer(s, 10)
    notes(s, "If they say VirusTotal already exists: VT is after-click research. We stop the navigation first and explain in plain English.")

    # 11 architecture
    s = new_slide(prs)
    add_text(s, 0.55, 0.28, 12, 0.25, ["10  /  HOW IT ACTUALLY RUNS"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.58, 12, 0.4, ["Input → normalize → intel + heuristics → fuse → explain → act"], size=20, bold=True)
    stack = [
        ("Frontend", "React 19, Vite, Tailwind, Axios, Recharts"),
        ("API", "FastAPI + Uvicorn. Not Flask."),
        ("Data", "SQLite threats.db. Not live Postgres."),
        ("Auth", "bcrypt + JWT"),
        ("Detect", "Heuristics + public feeds + weighted fusion"),
        ("QR", "OpenCV headless. Does not visit the URL."),
        ("Intel", "OpenPhish, URLhaus, PhishTank, Phishing Army, optional GSB"),
        ("Levels", "LOW 0–29 · MED 30–49 · HIGH 50–69 · CRIT 70–100"),
    ]
    for i, (title, body) in enumerate(stack):
        col, row = i % 4, i // 4
        card(s, 0.45 + col * 3.2, 1.2 + row * 2.7, 3.05, 2.5, title, [body], "")
    footer(s, 11)
    notes(s, "Fusion key ml is a copy of the lexical URL score — not sklearn. Psychology cannot solo-CRITICAL (cap 35).")

    # 12 close + demo
    s = new_slide(prs)
    add_text(s, 0.55, 0.28, 12, 0.25, ["11  /  DEMO  ·  HONEST LIMITS  ·  ASK"], size=12, bold=True, color=RED)
    add_text(s, 0.55, 0.58, 12, 0.45, ["We will show it. Then we will tell you what it is not."], size=22, bold=True)
    card(s, 0.55, 1.25, 6.1, 3.35, "Live demo path", [
        "1. Wake API: phisheye-api.onrender.com/health",
        "2. Dashboard: scan google.com → LOW",
        "3. Extension: github.com → gate, then open if LOW",
        "4. Lookalike stays closed. Do not open a live kit.",
        "",
        "Web  phisheye-web-fhn5.onrender.com",
        "Ext  /extension   API  /docs",
    ], "Show")
    card(s, 6.85, 1.25, 5.95, 3.35, "We do not claim", [
        "Flask  ·  live PostgreSQL  ·  trained CNN",
        "99% accuracy or published F1",
        "Reading WhatsApp, SMS, or Gmail inboxes",
        "Chrome Web Store listing",
        "Blockchain (theme includes it; we did not build it)",
        "Padlock = genuine bank",
    ], "Honesty")
    card(s, 0.55, 4.75, 12.25, 1.7, "What we want next", [
        "A labelled report set → then a supervised model. Persistent DB. Optional enterprise feed keys. The gate stays.",
    ], "Roadmap")
    footer(s, 12)
    notes(
        s,
        "Close: working prototype, explainable engine, scan-before-open. Free Render sleeps after ~15 minutes — hit /health first. Load unpacked, not Web Store.",
    )

    prs.save(str(OUT))
    try:
        DOWNLOADS.write_bytes(OUT.read_bytes())
    except OSError:
        pass
    print("Wrote", OUT)
    if DOWNLOADS.exists():
        print("Copied", DOWNLOADS)


if __name__ == "__main__":
    build()
