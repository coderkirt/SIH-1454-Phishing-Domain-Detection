# -*- coding: utf-8 -*-
"""Generate PHISHEYE_LENOVO_HACKATHON.pptx from verified repo facts."""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / "docs" / "_leap_media"
OUT = ROOT / "PHISHEYE_LENOVO_HACKATHON.pptx"

BG = RGBColor(0x0A, 0x0A, 0x0C)
CARD = RGBColor(0x14, 0x14, 0x18)
CARD2 = RGBColor(0x1A, 0x1A, 0x20)
RED = RGBColor(0xE1, 0x06, 0x00)
WHITE = RGBColor(0xF4, 0xF4, 0xF5)
MUTED = RGBColor(0x9C, 0x9C, 0xA4)
DIM = RGBColor(0x6B, 0x6B, 0x74)
LINE = RGBColor(0x2C, 0x2C, 0x34)
W, H = Inches(13.333), Inches(7.5)
TOTAL = 12
DISP, SANS, MONO = "Segoe UI", "Calibri", "Consolas"


def _font(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs", "a:latin"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def run(p, text, size=14, bold=False, color=WHITE, font=SANS):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    _font(r, font)
    return r


def tb(slide, l, t, w, h, lines, size=14, bold=False, color=WHITE, font=SANS,
       align=PP_ALIGN.LEFT, anchor="t", after=4):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", anchor)
    except Exception:
        pass
    for i, line in enumerate(lines):
        chunks = str(line).split("\n")
        for j, chunk in enumerate(chunks):
            p = tf.paragraphs[0] if i == 0 and j == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(after)
            run(p, chunk, size=size, bold=bold, color=color, font=font)
    return box


def rect(slide, l, t, w, h, fill=CARD, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def paint(slide):
    bg = rect(slide, 0, 0, 13.333, 7.5, BG)
    bg.left, bg.top = Emu(0), Emu(0)
    rect(slide, 0, 0, 0.07, 7.5, RED)
    rect(slide, 0, 0, 13.333, 0.035, RED)


def kicker(slide, text):
    tb(slide, 0.5, 0.22, 10, 0.22, [text], size=11, bold=True, color=RED, font=MONO)


def footer(slide, n):
    tb(slide, 0.5, 7.18, 9, 0.22, ["PHISHEYE  ·  LENOVO HACKATHON  ·  SOFTWARE"], size=10, color=DIM, font=MONO)
    tb(slide, 11.3, 7.18, 1.55, 0.22, [f"{n:02d}/{TOTAL:02d}"], size=10, color=DIM, font=MONO, align=PP_ALIGN.RIGHT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def pic(slide, path, l, t, w, h, cover=False):
    path = Path(path)
    bw, bh = Inches(w), Inches(h)
    if cover:
        crop_path = path.with_name(path.stem + "_cover.jpg")
        with Image.open(path) as im:
            im = im.convert("RGB")
            iw, ih = im.size
            target = w / h
            src = iw / ih
            if src > target:
                nw = int(ih * target)
                left = (iw - nw) // 2
                im = im.crop((left, 0, left + nw, ih))
            else:
                nh = int(iw / target)
                top = (ih - nh) // 2
                im = im.crop((0, top, iw, top + nh))
            im.save(crop_path, "JPEG", quality=88)
        return slide.shapes.add_picture(str(crop_path), Inches(l), Inches(t), bw, bh)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(bw / iw, bh / ih)
    nw, nh = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(l) + (bw - nw) / 2, Inches(t) + (bh - nh) / 2, nw, nh)


def node(slide, l, t, w, h, title, sub="", accent=False):
    rect(slide, l, t, w, h, RED if accent else CARD, LINE)
    if accent:
        rect(slide, l, t, 0.06, h, WHITE)
    else:
        rect(slide, l, t, 0.06, h, RED)
    title_h = 0.32 if sub else max(0.32, h - 0.16)
    anchor = "t" if sub else "ctr"
    tb(slide, l + 0.14, t + 0.08, w - 0.22, title_h, [title], size=13, bold=True, color=WHITE, font=SANS, after=0, anchor=anchor)
    if sub:
        tb(slide, l + 0.14, t + 0.42, w - 0.22, h - 0.52, [sub], size=12, color=MUTED if not accent else RGBColor(0xFF, 0xE4, 0xE4), font=SANS, after=2)


def chevron(slide, l, t, w=0.28, h=0.22):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RED
    s.line.fill.background()
    return s


def new(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    paint(s)
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 title
    s = new(prs)
    tb(s, 0.55, 1.55, 12, 0.28, ["LENOVO HACKATHON  ·  SOFTWARE  ·  CYBERSECURITY"], size=13, bold=True, color=RED, font=MONO)
    tb(s, 0.5, 2.0, 12.3, 1.15, ["PHISHEYE"], size=66, bold=True, color=WHITE, font=DISP, after=0)
    tb(s, 0.55, 3.25, 12, 0.55, ["Scan the link before the fake page loads."], size=26, color=WHITE, font=SANS)
    tb(s, 0.55, 4.05, 11.5, 0.9, [
        "Explainable phishing and lookalike-domain detection.",
        "One FastAPI engine. A React console. A Chrome scan-before-open gate.",
    ], size=16, color=MUTED)
    rect(s, 0.55, 5.2, 1.8, 0.07, RED)
    tb(s, 0.55, 5.45, 12, 0.7, [
        "Working software prototype  ·  No hardware  ·  No trained CNN  ·  SQLite live",
    ], size=13, color=DIM, font=MONO)
    footer(s, 1)
    notes(s, "Open with the product name once. Dashboard + extension, same engine. Do not invent team members.")

    # 2 problem
    s = new(prs)
    kicker(s, "01  /  THE PROBLEM")
    tb(s, 0.5, 0.48, 12, 0.45, ["A padlock does not mean it is your bank."], size=26, bold=True, font=DISP)
    steps = [
        ("ATTACKER", "Fake KYC / OTP lure"),
        ("TRUST", "Fear, urgency, brand"),
        ("FAKE LINK", "HTTPS lookalike"),
        ("VICTIM", "Enters OTP / UPI"),
        ("LOSS", "Account or money"),
    ]
    x = 0.5
    for i, (title, sub) in enumerate(steps):
        node(s, x, 1.35, 2.15, 1.55, title, sub, accent=(i == 4))
        if i < 4:
            chevron(s, x + 2.18, 1.95)
        x += 2.55
    node(s, 0.5, 3.25, 6.0, 3.15, "What the victim sees",
         "SMS or WhatsApp: “Your SBI account will be blocked. Complete KYC in 10 minutes. Share OTP.”\n\nThe page looks official. HTTPS is present. The warning in the browser is technical. The money still leaves.")
    node(s, 6.7, 3.25, 6.1, 3.15, "Who we designed for",
         "First-time UPI users and people who will not parse an SSL mismatch.\n\nSimple View is written in OTP language, not certificate jargon.")
    footer(s, 2)
    notes(s, "Tell the OTP story slowly. HTTPS is cheap. Point at LOSS.")

    # 3 gap
    s = new(prs)
    kicker(s, "02  /  THE GAP")
    tb(s, 0.5, 0.48, 12, 0.4, ["Feeds are late. Certificates are not brands. Panic is a signal."], size=22, bold=True, font=DISP)
    gaps = [
        ("Technical signals", "Lookalike host, new domain, TLS facts, SSRF-safe unwrap. Absence of a feed is not a clean bill of health."),
        ("Social engineering", "Urgency, fear, OTP pressure — English + Hinglish phrase rules, not a language model."),
        ("Threat intelligence", "A confirmed hit floors the score. Heuristics cannot talk it down. A miss is not SAFE."),
        ("Explainable output", "Simple View for humans who will not read a certificate. Technical View for operators."),
        ("In-browser act", "Scan first. LOW opens. MEDIUM asks. HIGH / CRITICAL stay closed."),
        ("Honesty bound", "No competitor ranking. No published F1. No trained CNN in this repository."),
    ]
    for i, (t, b) in enumerate(gaps):
        col, row = i % 3, i // 3
        node(s, 0.5 + col * 4.2, 1.15 + row * 2.75, 4.0, 2.55, t, b)
    footer(s, 3)
    notes(s, "VirusTotal is after-click research. We gate navigation and explain in plain English.")

    # 4 introducing
    s = new(prs)
    kicker(s, "03  /  PRODUCT")
    tb(s, 0.5, 0.48, 12, 0.4, ["PHISHEYE scores the lure. Then it tells the user what to do."], size=22, bold=True, font=DISP)
    flow = [
        ("INPUT", "URL · paste · QR · screenshot"),
        ("ANALYZE", "Normalize + signals"),
        ("FUSE / FLOOR", "Weights or intel floors"),
        ("RESULT", "LOW → CRITICAL"),
        ("ACT", "Explain · store · gate"),
    ]
    x = 0.45
    for i, (t, b) in enumerate(flow):
        node(s, x, 1.15, 2.2, 1.45, t, b, accent=(i == 4))
        if i < 4:
            chevron(s, x + 2.23, 1.7)
        x += 2.56
    node(s, 0.5, 2.9, 4.0, 3.5, "Protects", "People who will click a KYC SMS.\nOperators on the React console.\nAnyone using the Chrome gate.")
    node(s, 4.7, 2.9, 4.0, 3.5, "Analyzes (implemented)", "Direct URLs\nPasted SMS / WhatsApp / email text\nQR images (OpenCV, no visit)\nScreenshots (QR; OCR if Tesseract exists)\nVisible page links via extension")
    node(s, 8.9, 2.9, 3.9, 3.5, "Produces", "Risk 0–100 and a level\nSimple View warning\nTechnical evidence\nSQLite history\nExtension allow / hold / block")
    footer(s, 4)
    notes(s, "Only implemented channels. No inbox OAuth. QR does not fetch the phishing page.")

    # 5 architecture
    s = new(prs)
    kicker(s, "04  /  ARCHITECTURE  ·  FROM THE CODE")
    tb(s, 0.5, 0.48, 12, 0.35, ["Clients → FastAPI → analyzers → decide → SQLite → explain."], size=20, bold=True, font=DISP)
    rows = [
        ("USER", "Paste a URL or message, or navigate in Chrome.", "Two clients, one brain."),
        ("API", "FastAPI + Uvicorn on port 8000. Not Flask.", "JWT on mutating routes. GET /health."),
        ("SERVICES", "URL, brand, TLS, SSRF unwrap, QR (no visit).", "Intel feeds + psychology rules."),
        ("DECIDE", "URL path: intel floors cannot be averaged away.", "Content path: documented fusion weights."),
        ("STORE", "Live database is sqlite3 threats.db.", "Simple View + Technical View out."),
    ]
    y = 1.05
    for title, a, b in rows:
        node(s, 0.5, y, 2.2, 1.05, title, "", accent=True)
        node(s, 2.85, y, 5.0, 1.05, a)
        node(s, 8.0, y, 4.8, 1.05, b)
        y += 1.12
    footer(s, 5)
    notes(s, "URL scans: decision_engine. Content scans: fuse(). SQLAlchemy models are unused. Live DB is sqlite3.")

    # 6 AI/ML
    s = new(prs)
    kicker(s, "05  /  INTELLIGENCE  ·  WHAT “AI/ML” MEANS HERE")
    tb(s, 0.5, 0.48, 12, 0.4, ["Explainable scoring. No trained model file in this repository."], size=20, bold=True, font=DISP)
    pipe = [
        ("RAW INPUT", "URL or pasted text"),
        ("PREPROCESS", "Normalize · SSRF unwrap"),
        ("FEATURES", "TLD · lookalike · age · keywords"),
        ("SIGNALS", "Intel · psych · brand · community"),
        ("DECISION", "Floors + weighted fuse"),
    ]
    x = 0.4
    for i, (t, b) in enumerate(pipe):
        node(s, x, 1.1, 2.3, 1.5, t, b, accent=(i == 4))
        if i < 4:
            chevron(s, x + 2.33, 1.7)
        x += 2.58
    node(s, 0.5, 2.9, 6.05, 3.5, "What exists",
         "Heuristic URL features (TLD, lookalike, age, keywords)\nPhrase matching for panic / OTP language\nDocumented fusion weights — starting defaults, not proven-optimal\nFusion key “ml” copies the URL score so a low ML number cannot prove safety\nNumPy reshapes images for OpenCV — not neural nets")
    node(s, 6.75, 2.9, 6.05, 3.5, "What does not exist",
         "No sklearn / TensorFlow / CNN / BERT\nNo labelled production dataset\nNo published accuracy, precision, recall, or F1\nmodel_confidence = how many signals ≥ 20 — not accuracy\nSupervised model = FUTURE after curated reports")
    footer(s, 6)
    notes(s, "If they ask for CNN accuracy: there is none. Phase 2 after labels.")

    # 7 psychology
    s = new(prs)
    kicker(s, "06  /  PSYCHOLOGICAL THREAT DETECTION")
    tb(s, 0.5, 0.48, 12, 0.4, ["Rules that read panic. Not a neural language model."], size=22, bold=True, font=DISP)
    node(s, 0.5, 1.1, 6.1, 2.35, "Example (implemented patterns)",
         "“Your account will be blocked. Verify immediately.”\n\nHits: fear_tactics · urgency_detected · credential_pressure\nEnglish + Hinglish: abhi click, band ho jayega, enter otp, complete kyc")
    node(s, 6.8, 1.1, 6.0, 2.35, "Guardrail",
         "Psychology cannot take a site to HIGH/CRITICAL alone.\nIf no strong URL / domain / intel / brand / community signal, fuse caps the score at 35 (MEDIUM).")
    cats = [
        ("Urgency", "act now · jaldi karo"),
        ("Fear", "account blocked"),
        ("Impersonation", "RBI · cyber cell"),
        ("Coercion", "mandatory KYC"),
        ("Money pressure", "pay immediately"),
        ("OTP / PIN", "enter otp · share otp"),
        ("Countdown", "within 10 minutes"),
        ("Reward bait", "you have won"),
        ("Secrecy", "do not tell anyone"),
        ("Emotional ask", "help me urgently"),
    ]
    for i, (c, ex) in enumerate(cats):
        col, row = i % 5, i // 5
        node(s, 0.5 + col * 2.52, 3.65 + row * 1.5, 2.4, 1.35, c, ex)
    footer(s, 7)
    notes(s, "Read the example. Then the cap. Then Hinglish. Do not invent a psych score for this example.")

    # 8 channels + extension
    s = new(prs)
    kicker(s, "07  /  CHANNELS  ·  BROWSER GATE")
    tb(s, 0.5, 0.46, 7, 0.35, ["Scan first. Then maybe open."], size=22, bold=True, font=DISP)
    chans = [
        ("URL", "check-url"),
        ("Pasted text", "SMS / WA / email paste"),
        ("QR", "OpenCV · no visit"),
        ("Screenshot", "QR + optional OCR"),
        ("Extension", "Visible links + navigation"),
    ]
    for i, (t, b) in enumerate(chans):
        node(s, 0.5, 0.95 + i * 1.1, 4.35, 1.0, t, b)
    node(s, 5.05, 0.95, 3.55, 5.4, "Extension flow",
         "User navigates\n→ service worker holds tab\n→ POST check-url\n→ LOW: auto-open\n→ MEDIUM: ask to continue\n→ HIGH/CRITICAL: stay closed\n\nMV3 1.8.0\nLoad unpacked — not Web Store\nDoes not read passwords or SMS")
    rect(s, 8.75, 0.95, 4.1, 5.4, CARD, LINE)
    if (MEDIA / "13_extension.png").exists():
        pic(s, MEDIA / "13_extension.png", 8.85, 1.05, 3.9, 5.2, cover=True)
    footer(s, 8)
    notes(s, "Demo github.com LOW. Lookalike stays closed. Do not open a live kit.")

    # 9 risk
    s = new(prs)
    kicker(s, "08  /  RISK  ·  EXPLAINABILITY")
    tb(s, 0.5, 0.48, 12, 0.35, ["Two views. One score. Documented combination — not a secret formula."], size=20, bold=True, font=DISP)
    node(s, 0.5, 1.05, 6.1, 2.7, "SIMPLE VIEW",
         "This website may be pretending to be a trusted service.\nDo not enter your password, OTP, or card number.\n\nVerdicts: NO MALICIOUS INDICATORS · BE CAREFUL · DANGER")
    node(s, 6.8, 1.05, 6.0, 2.7, "TECHNICAL VIEW",
         "Domain / lookalike tags · TLS · RDAP age\nThreat-intel provider + match type\nHeuristic reasons · psych findings\nmodel_confidence = signal agreement")
    node(s, 0.5, 3.95, 6.1, 2.5, "URL decision (floors)",
         "Confirmed intel: max(heuristic, floor)\nexact URL 85 · normalized 80\nhostname 72 · registered domain 55\nHeuristics cannot veto a confirmed hit.")
    node(s, 6.8, 3.95, 6.0, 2.5, "Content fusion (weights)",
         "url 0.22 · domain 0.16 · ml 0.10 · nlp 0.10\npsych 0.10 · intel 0.10 · brand 0.08\nredirect 0.04 · sender 0.05 · community 0.05\nLOW 0–29 · MED 30–49 · HIGH 50–69 · CRIT 70–100")
    footer(s, 9)
    notes(s, "Quote the weights from the slide — they are in risk_engine.py. Confidence is not accuracy.")

    # 10 stack
    s = new(prs)
    kicker(s, "09  /  STACK  ·  ACTUALLY USED")
    tb(s, 0.5, 0.48, 12, 0.35, ["Names plus one job. No logo wall. No Postgres."], size=20, bold=True, font=DISP)
    groups = [
        ("FRONTEND", "React 19 dashboard\nVite 8 + Tailwind 4\nAxios to FastAPI\nRecharts on Statistics"),
        ("BACKEND", "FastAPI + Uvicorn\nPydantic request models\nbcrypt passwords\nJWT on mutating routes"),
        ("DATA", "Live: sqlite3 threats.db\nusers · url_checks · scans\nPostgres = FUTURE, not live"),
        ("INTELLIGENCE", "Heuristics + fusion\nPublic threat feeds\nNot a trained network\nml channel copies URL score"),
        ("SECURITY", "SSRF before unwrap\nParameterized SQL\n5 MB upload cap\nExtension does not read passwords"),
        ("EXTENSION", "Chrome MV3 1.8.0\nScan-before-open gate\nSame check-url API\nLoad unpacked, not Web Store"),
        ("EXTERNAL", "OpenPhish · URLhaus\nPhishTank · Phishing Army\nOptional Google Safe Browsing\nMiss ≠ SAFE"),
        ("DEPLOY", "Render FastAPI service\nStatic React site\nExtension zip on /extension\nFree API sleeps ~15 min"),
    ]
    for i, (t, b) in enumerate(groups):
        col, row = i % 4, i // 4
        node(s, 0.45 + col * 3.2, 1.1 + row * 2.8, 3.05, 2.55, t, b)
    footer(s, 10)
    notes(s, "If they say AI, point at fusion + intel. If they say database, say SQLite.")

    # 11 product
    s = new(prs)
    kicker(s, "10  /  WORKING PRODUCT")
    tb(s, 0.5, 0.46, 12, 0.3, ["Screenshots of the running UI. Dashboard counts are not an F1 score."], size=16, bold=True, font=DISP)
    shots = [
        (MEDIA / "01_landing.png", 0.4, 0.88, 6.2, 3.0),
        (MEDIA / "05_scan.png", 6.75, 0.88, 6.15, 3.0),
        (MEDIA / "07_history.png", 0.4, 4.0, 6.2, 2.95),
        (MEDIA / "02_dashboard.png", 6.75, 4.0, 6.15, 2.95),
    ]
    for path, l, t, w, h in shots:
        rect(s, l, t, w, h, CARD, LINE)
        if path.exists():
            pic(s, path, l + 0.08, t + 0.08, w - 0.16, h - 0.16, cover=True)
    footer(s, 11)
    notes(s, "These are real screens. Do not quote 87.8% as accuracy. It is a dashboard ratio on that SQLite file.")

    # 12 impact + future
    s = new(prs)
    kicker(s, "11  /  IMPACT  ·  NEXT  ·  CLOSE")
    tb(s, 0.5, 0.48, 12, 0.4, ["Protect people who cannot read a certificate warning."], size=22, bold=True, font=DISP)
    node(s, 0.5, 1.1, 6.1, 3.2, "Impact now (implemented)",
         "Safer click path: scan before open\nPlain-English OTP warnings\nMulti-channel paste analysis\nCommunity labels as a signal, not proof")
    node(s, 6.8, 1.1, 6.0, 3.2, "FUTURE / NEXT PHASE",
         "Curated reports → supervised model\nPostgreSQL at scale\nAndroid SMS companion\nOfficial inbox APIs if granted\n\nNot on this slide as current work.")
    rect(s, 0.5, 4.5, 12.3, 2.0, CARD, LINE)
    rect(s, 0.5, 4.5, 0.08, 2.0, RED)
    tb(s, 0.8, 4.7, 11.7, 1.6, [
        "Phishing + social engineering + threat intelligence + explainable scoring + browser protection.",
        "That combination is PHISHEYE.",
    ], size=20, bold=True, color=WHITE, font=DISP)
    footer(s, 12)
    notes(s, "Say FUTURE out loud. Close on the last line. No hardware. No blockchain.")

    prs.save(str(OUT))
    print("Wrote", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    build()
