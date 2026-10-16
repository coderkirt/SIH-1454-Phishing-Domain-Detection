"""
Build PhishEye_Synopsis_v2.docx from the official CSE mini-project template.

Option B wording: the prototype is a hybrid, rule- and feed-driven detection
system; a supervised ML classifier is future scope. No project source file is
modified. Facts follow the live implementation (FastAPI + sqlite3 + React +
Chrome MV3); no metrics or hardware are invented.

Figures come from synopsis_generator/figures/ (render_figures.py and
capture_screenshots.py). Run those first, or let this script render the
diagrams if they are missing.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
GENERATOR = Path(__file__).resolve().parent
FIGDIR = GENERATOR / "figures"
TEMPLATE_CANDIDATES = [
    Path(r"c:\Users\HRITIK KUMAR\Downloads\CSE_Mini Project Synopsis_Format 2026_27.docx"),
    ROOT / "CSE_Mini Project Synopsis_Format 2026_27.docx",
    GENERATOR / "CSE_Mini Project Synopsis_Format 2026_27.docx",
]
OUTPUT = ROOT / "PhishEye_Synopsis_v2.docx"
FONT = "Times New Roman"
ACCENT = "0B1F3A"          # navy header fill
BORDER = "BFBFBF"          # light grey borders
TEXT_W_CM = 21.0 - 2 * 2.54  # usable text width on A4 with 2.54 cm margins

TITLE = "PhishEye: AI & ML based Cyberguard"
PROJECT_ID = "26_CS_2L_19"
STUDENTS = ["Hritik Kumar (2501640100427)", "Kirt Raj Dixit (2501640100468)"]
SUPERVISOR = "Prasun Tripathi"
DESIGNATION = "Assistant Professor"
ACCESSED = "Accessed: Oct. 3, 2026"

# Height caps (cm) per figure, tuned so that figures share pages with the text
# that introduces them and the document stays within 12 pages.
FIG_HEIGHTS = {
    "fig1_system_architecture": 9.6,
    "fig2_dfd_level0": 6.4,
    "fig3_dfd_level1": 8.0,
    "fig4_use_case": 8.0,
    "fig5_er_diagram": 9.2,
    "fig6_risk_decision_flow": 8.6,
    "fig7_extension_gating_flow": 7.8,
    "fig8_scan_result": 8.2,
    "fig9_history": 5.8,
}


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #
def _set_run_font(run, size_pt: float, *, bold=False, italic=False, color="000000"):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _set_paragraph_format(paragraph, *, size=12, bold=False, align="justify", space_after=6, space_before=0,
                          line="1.5", keep_next=False):
    pf = paragraph.paragraph_format
    if line == "1.5":
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.keep_with_next = keep_next
    paragraph.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    for run in paragraph.runs:
        _set_run_font(run, size, bold=bold or bool(run.bold), italic=bool(run.italic))


def _clear_runs(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def _replace_paragraph_text(paragraph, text, *, size=None, bold=None, keep_size=True):
    first_size = paragraph.runs[0].font.size if paragraph.runs else None
    first_bold = paragraph.runs[0].bold if paragraph.runs else None
    _clear_runs(paragraph)
    run = paragraph.add_run(text)
    _set_run_font(run, 12)
    if keep_size and first_size is not None:
        run.font.size = first_size
    elif size is not None:
        run.font.size = Pt(size)
    run.bold = first_bold if bold is None else bold
    return run


def _delete_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _insert_paragraph_after(paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._element.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def _add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    _set_run_font(run, 11)
    r = run._r
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "2"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text is not None:
            el.text = text
        r.append(el)


def _find_template() -> Path:
    for path in TEMPLATE_CANDIDATES:
        if path.exists():
            return path
    raise FileNotFoundError("CSE mini-project synopsis template not found (Downloads or project root).")


def _is_intro_heading(text: str) -> bool:
    return text.strip().lower().replace("\t", "").startswith("introduction")


# --------------------------------------------------------------------------- #
# cover page
# --------------------------------------------------------------------------- #
def _scale_cover_logo(doc: Document, factor: float):
    ns_wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    ns_a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for para in doc.paragraphs:
        if _is_intro_heading(para.text):
            break
        for extent in para._p.findall(f".//{ns_wp}extent"):
            extent.set("cx", str(int(int(extent.get("cx")) * factor)))
            extent.set("cy", str(int(int(extent.get("cy")) * factor)))
        for ext in para._p.findall(f".//{ns_a}ext"):
            if ext.get("cx") and ext.get("cy"):
                ext.set("cx", str(int(int(ext.get("cx")) * factor)))
                ext.set("cy", str(int(int(ext.get("cy")) * factor)))


def prepare_cover(doc: Document):
    student_paras, address_paras = [], []
    for para in doc.paragraphs:
        raw = para.text
        if "TITLE OF THE PROJECT" in raw.upper():
            _replace_paragraph_text(para, TITLE, bold=True)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif "Project Id" in raw:
            _replace_paragraph_text(para, f"(Project Id: {PROJECT_ID})", bold=True)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif raw.strip() == "Student Name (Roll No)":
            student_paras.append(para)
        elif raw.strip() == "Supervisor Name":
            _replace_paragraph_text(para, SUPERVISOR, bold=True)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif raw.strip() == "(Designation)":
            _replace_paragraph_text(para, f"({DESIGNATION})", bold=False)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif raw.strip().startswith("Pranveer Singh Institute"):
            _clear_runs(para)
            r1 = para.add_run("Pranveer Singh Institute of Technology")
            _set_run_font(r1, 14, bold=True)
            r1.add_break()
            r2 = para.add_run("Kanpur - Agra - Delhi National Highway - 19, Bhauti - Kanpur - 209305")
            _set_run_font(r2, 12)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.15
        elif "National Highway" in raw or ("Bhauti" in raw and "209305" in raw):
            address_paras.append(para)

    for para, name in zip(student_paras[:2], STUDENTS):
        para.style = doc.styles["Normal"]
        _replace_paragraph_text(para, name, size=14, bold=True, keep_size=False)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.15
    for para in student_paras[2:]:
        _delete_paragraph(para)
    for para in address_paras:
        _delete_paragraph(para)

    _scale_cover_logo(doc, 0.80)

    # Normalise spacing on the cover; keep one blank line between blocks so the
    # page stays visually balanced on A4 (the template was US Letter).
    kept_blank = 0
    for para in list(doc.paragraphs):
        if _is_intro_heading(para.text):
            break
        has_drawing = bool(para._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"))
        if para.text.strip() == "" and not has_drawing:
            if kept_blank >= 1:
                _delete_paragraph(para)
                continue
            kept_blank += 1
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.6
        elif para.text.strip():
            kept_blank = 0
            para.paragraph_format.line_spacing = 1.15
            if para.paragraph_format.space_before is None or para.paragraph_format.space_before > Pt(10):
                para.paragraph_format.space_before = Pt(6)
        else:
            para.paragraph_format.space_before = Pt(10)
            para.paragraph_format.space_after = Pt(10)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def strip_placeholders(doc: Document) -> Paragraph:
    """Remove the template body (from the Introduction heading to the end). Return the last cover paragraph."""
    paras = list(doc.paragraphs)
    start = next((p for p in paras if _is_intro_heading(p.text)), None)
    if start is None:
        raise RuntimeError("Could not find the Introduction heading in the template.")
    last_cover = None
    for para in paras:
        if para._element is start._element:
            break
        last_cover = para
    remove = False
    for para in paras:
        if para._element is start._element:
            remove = True
        if remove:
            _delete_paragraph(para)
    return last_cover


# --------------------------------------------------------------------------- #
# body builder
# --------------------------------------------------------------------------- #
class Builder:
    def __init__(self, doc: Document, last: Paragraph):
        self.doc = doc
        self.cursor = last
        self.fig_n = 0
        self.tab_n = 0

    # paragraphs -------------------------------------------------------------
    def add(self, text="", *, size=12, bold=False, italic=False, align="justify", space_after=4,
            space_before=0, line="1.5", keep_next=False):
        p = _insert_paragraph_after(self.cursor)
        self.cursor = p
        if text:
            run = p.add_run(text)
            _set_run_font(run, size, bold=bold, italic=italic)
        _set_paragraph_format(p, size=size, bold=bold, align=align, space_after=space_after,
                              space_before=space_before, line=line, keep_next=keep_next)
        return p

    def heading(self, text, level=1):
        size = 14 if level == 1 else 12
        return self.add(text, size=size, bold=True, align="left", space_after=4,
                        space_before=10 if level == 1 else 6, keep_next=True)

    def rich(self, parts, **kw):
        """parts = [(text, bold), ...]"""
        p = self.add("", **kw)
        for text, bold in parts:
            r = p.add_run(text)
            _set_run_font(r, kw.get("size", 12), bold=bold)
        return p

    def numbered(self, items, prefix="", size=12):
        for i, item in enumerate(items, 1):
            p = self.add("", size=size, space_after=4)
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            r = p.add_run(f"{prefix}{i}.\t")
            _set_run_font(r, size, bold=True)
            r2 = p.add_run(item)
            _set_run_font(r2, size)
            # tab stop at indent
            tabs = p.paragraph_format.tab_stops
            tabs.add_tab_stop(Cm(0.9))

    def page_break(self):
        p = self.add("", space_after=0, line="single")
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    # figures -----------------------------------------------------------------
    def figure(self, path: Path, caption: str, max_w_cm=15.5, max_h_cm=15.0):
        max_h_cm = FIG_HEIGHTS.get(path.stem, max_h_cm)
        self.fig_n += 1
        with Image.open(path) as im:
            w, h = im.size
        width = min(max_w_cm, max_h_cm * w / h)
        p = self.add("", align="center", space_after=2, space_before=2, line="single", keep_next=True)
        p.add_run().add_picture(str(path), width=Cm(width))
        cap = self.add(f"Fig. {self.fig_n}: {caption}", size=11, italic=True, align="center",
                       space_after=6, line="single")
        return cap

    # tables ------------------------------------------------------------------
    def table(self, caption: str, headers, rows, widths_cm, *, size=10.5, center_cols=(), bold_first_col=False):
        self.tab_n += 1
        cap = self.add(f"Table {self.tab_n}: {caption}", size=11, bold=True, align="left",
                       space_after=3, space_before=4, line="single", keep_next=True)

        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        self.cursor._p.addnext(tbl._tbl)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        # borders + layout
        tblPr = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), BORDER)
            borders.append(el)
        tblPr.append(borders)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        cellmar = OxmlElement("w:tblCellMar")
        for side, val in (("top", 40), ("bottom", 40), ("left", 90), ("right", 90)):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(val))
            el.set(qn("w:type"), "dxa")
            cellmar.append(el)
        tblPr.append(cellmar)

        def fill(cell, text, *, bold=False, color="000000", shade=None, align="left"):
            p = cell.paragraphs[0]
            _clear_runs(p)
            lines = text.split("\n")
            for i, line in enumerate(lines):
                r = p.add_run(line)
                _set_run_font(r, size, bold=bold, color=color)
                if i < len(lines) - 1:
                    r.add_break()
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
            if shade:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), shade)
                tcPr.append(shd)

        # header row
        hdr = tbl.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        trPr.append(th)
        for j, text in enumerate(headers):
            fill(hdr.cells[j], text, bold=True, color="FFFFFF", shade=ACCENT,
                 align="center" if j in center_cols else "left")
        for i, row in enumerate(rows, 1):
            tr = tbl.rows[i]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            for j, text in enumerate(row):
                fill(tr.cells[j], str(text), bold=(bold_first_col and j == 0),
                     align="center" if j in center_cols else "left")
        for row in tbl.rows:
            for j, wcm in enumerate(widths_cm):
                row.cells[j].width = Cm(wcm)
        # column grid
        for j, gc in enumerate(tbl._tbl.tblGrid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths_cm[j] * 567)))

        # spacer paragraph after table; cursor moves there
        spacer = OxmlElement("w:p")
        tbl._tbl.addnext(spacer)
        self.cursor = Paragraph(spacer, self.cursor._parent)
        _set_paragraph_format(self.cursor, size=4, space_after=2, line="single", align="left")
        return tbl


# --------------------------------------------------------------------------- #
# content
# --------------------------------------------------------------------------- #
def write_body(doc: Document, last_cover: Paragraph, fig: dict):
    b = Builder(doc, last_cover)
    br = _insert_paragraph_after(b.cursor)
    b.cursor = br
    br.add_run().add_break(WD_BREAK.PAGE)
    _set_paragraph_format(br, space_after=0, line="single", align="left")

    # ------------------------------------------------------------------ 1
    b.heading("1. Introduction")
    b.add(
        "PhishEye: AI & ML based Cyberguard is a hybrid, AI-assisted phishing detection system that helps a "
        "person decide whether a link or a message is safe before acting on it. Phishing is the most common "
        "entry point for online fraud: a message on SMS, WhatsApp or e-mail imitates a bank, a courier or a "
        "government office and pushes the reader to open a link and type a password, an OTP, a UPI PIN or card "
        "details. The fake page often shows a valid padlock, a familiar logo and an address that differs from "
        "the real one by a single letter, so a non-technical user cannot tell it apart. Browser blocklists react "
        "only after a site is reported and say nothing about the wording that lured the user there."
    )
    b.add(
        "PhishEye addresses this gap with a software-only system made of a web console, a REST service, a local "
        "database and a Chrome extension. The user submits a web address, pastes a suspicious message, or uploads "
        "a QR-code image. Three analysis engines run on the same input: URL and domain heuristics (lookalike "
        "brand names, newly registered domains, suspicious endings, missing HTTPS), live lookups in public "
        "threat-intelligence feeds, and natural-language rules that detect psychological pressure such as "
        "urgency, fear, authority and OTP or password requests, in English and Hinglish. A weighted fusion step "
        "combines the signals into a 0-100 risk score and one of four levels (LOW, MEDIUM, HIGH, CRITICAL). The "
        "result is explained in a Simple View for ordinary users and a Technical View for operators, and the "
        "extension applies the same decision to block a dangerous page before it opens. The current prototype is "
        "rule- and feed-driven; training a supervised machine-learning classifier on labelled phishing data is "
        "planned as future scope (Section 4.8)."
    )
    b.add(
        "Key terms: phishing - impersonation of a trusted party to steal credentials or money; lookalike domain - "
        "an address crafted to resemble a genuine brand; threat intelligence - publicly shared lists of confirmed "
        "malicious addresses; psychological manipulation cues - wording that pressures the reader to act without "
        "checking; explainability - presenting the reasons behind a verdict in plain language; risk level - the "
        "category assigned to a score so that a non-expert can act on it."
    )

    # ------------------------------------------------------------------ 2
    b.heading("2. Project Objective")
    b.add("The project sets out to achieve the following outcomes:")
    b.numbered([
        "To detect phishing and lookalike-domain web addresses, whether typed directly, pasted inside a message, "
        "or encoded in a QR image, and to express the result as a 0-100 risk score with a named level.",
        "To identify social-engineering pressure in message text - urgency, fear, authority, prize bait and "
        "requests for OTPs or passwords - in both English and Hinglish.",
        "To cross-check every address against public threat-intelligence feeds so that a confirmed phishing "
        "report always produces a high score, whatever the other signals say.",
        "To explain each verdict in two forms: a Simple View written for ordinary users and a Technical View "
        "listing the signals, tags and feed matches for operators.",
        "To protect users at the moment of risk through a Chrome extension that scans a link before the page "
        "opens and keeps high-risk or unverifiable destinations closed.",
        "To maintain a secure, per-user history of scans, community reports and feedback in a lightweight "
        "local database, with hashed passwords and token-based login.",
        "To design the scoring pipeline so that a supervised machine-learning classifier can be added later as "
        "one more signal without changing the user-facing workflow.",
    ])
    b.add(
        "The primary beneficiary is a student, a first-time digital-banking user, or anyone who receives OTP and "
        "KYC lures and does not read certificate details: they get a clear answer and a short reason before they "
        "click. Console operators gain a searchable history, statistics and community reports that reveal "
        "recurring scam domains. The system never reads private inboxes or phone SMS; it analyses only what the "
        "user submits or the page the user is about to visit.",
        space_before=4,
    )
    b.table(
        "Key features of PhishEye",
        ["Feature", "What it does for the user"],
        [
            ["Multi-input scanning", "Accepts a URL, pasted SMS/WhatsApp/e-mail text or a QR image; scores every link found"],
            ["URL and domain heuristics", "Flags lookalike brands, new domains, suspicious endings, shorteners, missing HTTPS"],
            ["Threat-intelligence lookup", "Checks OpenPhish, URLhaus, PhishTank, Phishing Army (Safe Browsing with a key)"],
            ["Psychological-cue detection", "Finds urgency, fear, authority, greed bait, OTP/password requests (English, Hinglish)"],
            ["Weighted risk fusion", "Combines all signals into a 0-100 score and a LOW / MEDIUM / HIGH / CRITICAL level"],
            ["Simple and Technical View", "Plain-language warning for everyone; detailed evidence for operators"],
            ["Chrome extension gate", "Scans a link before the page opens; blocks HIGH, CRITICAL and unverifiable pages"],
            ["History and community reports", "Stores scans per user; users can mark a domain scam, risky or safe as one extra signal"],
        ],
        [4.6, 11.3],
        bold_first_col=True,
    )

    # ------------------------------------------------------------------ 3
    b.heading("3. Feasibility Study")
    b.heading("3.1 Technical Feasibility", level=2)
    b.add(
        "The system is built from mature open-source components - Python 3.12 with FastAPI, SQLite, React 19 "
        "with Vite, OpenCV and the Chrome Manifest V3 platform - all already integrated and running. No graphics "
        "card or special server is needed because detection uses rules, lookups and arithmetic rather than a "
        "neural network. Threat feeds are reached over HTTPS with time-outs and a short local cache."
    )
    b.heading("3.2 Operational Feasibility", level=2)
    b.add(
        "A user registers, opens the Scan page, submits a link, message or image, and reads the Simple View "
        "first. The extension installs through Chrome's developer mode, so no store listing is needed. If the "
        "service cannot be reached, the extension keeps the page closed: an unknown page is not a safe page."
    )
    b.heading("3.3 Economic Feasibility", level=2)
    b.add(
        "Every library is free and open source and the database is a local file with no licence fee. The threat "
        "feeds are public; Google Safe Browsing is optional and off without a key. A free hosting tier suffices "
        "for demonstration and development runs on a student laptop, so there is no mandatory recurring cost."
    )
    b.heading("3.4 Schedule Feasibility", level=2)
    b.add(
        "The work is organised into six phases over a twelve-week semester plan (Table 5). The engines, scoring, "
        "console, extension and tests already exist as a working prototype, so the remaining effort concentrates "
        "on evaluation, documentation and future-scope items, with slack left for review feedback."
    )
    b.heading("3.5 Legal and Ethical Feasibility", level=2)
    b.add(
        "The system stores only an account (user name, e-mail, hashed password) and the addresses or results the "
        "user submits. It never captures passwords, OTPs or card numbers from web pages; the extension does not "
        "read form fields, cookies or phone messages; QR images are decoded in memory and not saved. The scanner "
        "is blocked from visiting internal network addresses, and community reports are shown as user opinions."
    )
    b.table(
        "Feasibility summary",
        ["Type", "Assessment"],
        [
            ["Technical", "Feasible - proven open-source stack, already integrated; no GPU needed"],
            ["Operational", "Feasible - simple web workflow; fails closed when the service is unreachable"],
            ["Economic", "Feasible - no licence or mandatory hosting cost; paid feed optional"],
            ["Schedule", "Feasible - prototype complete; twelve-week plan with slack for write-up"],
            ["Legal / ethical", "Feasible - minimal personal data, hashed passwords, no inbox access"],
        ],
        [3.2, 12.7],
        bold_first_col=True,
    )

    # ------------------------------------------------------------------ 4
    b.heading("4. Methodology / Planning of Work")
    b.add(
        "The project follows an incremental, phase-wise methodology; each phase produces a reviewable artefact, "
        "and the figures and tables below follow the phase they describe."
    )

    b.heading("4.1 Phase 1 - Problem Study and Requirement Analysis", level=2)
    b.add(
        "Real phishing lures (bank KYC, courier fee, electricity disconnection, prize and job scams) were "
        "collected from public awareness material and studied for technical tricks in the address and "
        "psychological tricks in the wording. The study produced the requirements in Section 2, the four-level "
        "risk scale and the rule that an unverifiable result must never be reported as safe."
    )

    b.heading("4.2 Phase 2 - System Design", level=2)
    b.add(
        "The architecture is layered (Fig. 1). Clients talk to a single REST service, which runs the analysis "
        "engines, hands their signals to the risk-decision engine, builds the two explanation views and stores "
        "results in SQLite. The QR decoder feeds the decoded address into the same URL analysis path as a typed link."
    )
    b.figure(fig["fig1_system_architecture"], "Layered system architecture of PhishEye.", max_h_cm=11.0)
    b.add(
        "Fig. 2 shows the system as one process exchanging data with the user, the extension and the feed "
        "services; Fig. 3 expands it into six numbered processes and two data stores, D1 (SQLite database) and "
        "D2 (local feed cache)."
    )
    b.figure(fig["fig2_dfd_level0"], "Data-flow diagram, Level 0 (context diagram).", max_h_cm=6.6)
    b.figure(fig["fig3_dfd_level1"], "Data-flow diagram, Level 1 with numbered processes and data stores.", max_h_cm=9.0)
    b.add(
        "The use-case view (Fig. 4) identifies three actors. A visitor can scan a link, message or QR image "
        "without logging in; viewing history, reporting a domain and giving feedback require a signed-in "
        "account. The extension acts as an automated client; the feed services are external systems consulted "
        "during a scan."
    )
    b.figure(fig["fig4_use_case"], "Use-case diagram of PhishEye.", max_h_cm=9.0)
    b.add(
        "The database design (Fig. 5) keeps one table per concern. Solid lines mark relationships enforced by a "
        "declared foreign key; dashed lines mark references stored as plain identifier columns and maintained by "
        "the application. The domain-reputation table is independent and keyed by domain name."
    )
    b.figure(fig["fig5_er_diagram"], "Entity-relationship diagram of the SQLite database (crow's-foot notation).", max_h_cm=11.0)

    b.heading("4.3 Phase 3 - Implementation of the Analysis Engines", level=2)
    b.add(
        "Three engines were implemented and tested independently. The URL and domain engine normalises the "
        "address, safely expands shortened links, checks whether the domain exists and how recently it was "
        "registered (under 30 days adds strongly to the score), inspects the certificate, and matches the host "
        "against brand and lookalike patterns. The threat-intelligence engine queries each feed with a time-out "
        "and records match, no match or unavailable, so a feed outage is visible rather than treated as clean. "
        "The psychology engine applies phrase rules in ten categories (urgency, fear, authority, coercion, "
        "financial pressure, credential pressure, countdown, greed bait, secrecy and emotional manipulation), "
        "each with a fixed point value; wording alone can raise a score to at most 35, so a harmless message "
        "with urgent words is never blocked outright."
    )

    b.heading("4.4 Phase 4 - Risk Decision and Explanation", level=2)
    b.add(
        "For content with several links, the engine signals are combined with the fixed weights in Table 3 "
        "(they sum to 1.0). For a single address, the heuristic score is compared with a minimum score tied to "
        "the strength of any feed match (exact address 85, normalised address 80, host name 72, registered "
        "domain 55) and the higher value wins, so heuristics can never cancel a confirmed report. The final score "
        "is mapped to a level (Table 4); the same decision drives both views and the extension (Fig. 6)."
    )
    b.table(
        "Fusion weights used to combine engine signals",
        ["Signal", "Weight", "What it measures"],
        [
            ["URL structure", "0.22", "Lexical tricks in the address: lookalike words, odd characters"],
            ["Domain reputation", "0.16", "Domain existence, age, suspicious ending, brand lookalike"],
            ["Lexical score", "0.10", "Compact numeric summary of the URL features"],
            ["Message language", "0.10", "Scam vocabulary and link-bait phrasing in the text"],
            ["Psychological cues", "0.10", "Urgency, fear, authority, OTP/password pressure"],
            ["Threat intelligence", "0.10", "Match status across the public feeds"],
            ["Brand impersonation", "0.08", "Named bank or brand referenced by a non-official domain"],
            ["Sender signals", "0.05", "Sender identity or header hints when supplied"],
            ["Community reports", "0.05", "Scam / risky / safe marks from other users"],
            ["Redirect behaviour", "0.04", "Number of redirect hops and shortener use"],
        ],
        [3.9, 1.8, 10.2],
        center_cols=(1,),
        bold_first_col=True,
    )
    b.table(
        "Risk levels and the action taken",
        ["Level", "Score range", "Meaning for the user", "Extension action"],
        [
            ["LOW", "0 - 29", "No malicious indicators found", "Page opens"],
            ["MEDIUM", "30 - 49", "Some suspicious signs; check before trusting", "User is asked to confirm"],
            ["HIGH", "50 - 69", "Strong phishing indicators", "Page is blocked"],
            ["CRITICAL", "70 - 100", "Confirmed or near-certain phishing", "Page is blocked"],
        ],
        [2.6, 2.6, 6.3, 4.4],
        center_cols=(1,),
        bold_first_col=True,
    )
    b.figure(fig["fig6_risk_decision_flow"], "Risk decision flow from input to colour-coded level.", max_h_cm=10.5)

    b.heading("4.5 Phase 5 - Browser Extension", level=2)
    b.add(
        "The extension intercepts navigation to a new address, holds the tab, and asks the service for a verdict "
        "(Fig. 7). LOW opens the page; MEDIUM shows the warning and lets the user decide; HIGH and CRITICAL stay "
        "blocked. If the service times out or is unreachable the page also stays blocked (\"unknown is not "
        "safe\"). Verdicts are cached for 90 seconds to avoid repeated scans of the same address."
    )
    b.figure(fig["fig7_extension_gating_flow"], "Chrome extension scan-before-open gating flow.", max_h_cm=9.5)

    b.heading("4.6 Phase 6 - Testing, Deployment and Documentation", level=2)
    b.add(
        "Automated tests cover the API routes, psychology rules, feed-status handling, QR decoding, certificate "
        "helpers and address normalisation, and a set of known phishing and legitimate samples is re-run after "
        "every change. The service and console are deployed on a free hosting tier for demonstration. Fig. 8 "
        "shows a real scan of a lookalike banking address; Fig. 9 shows the per-user history page."
    )
    b.figure(fig["fig8_scan_result"], "Scan result page: risk score, classification, recommendation and technical evidence.", max_h_cm=10.5)
    b.figure(fig["fig9_history"], "Report history page listing stored scans with score, level and status.", max_h_cm=6.5)

    b.heading("4.7 Work Schedule", level=2)
    b.table(
        "Phase-wise work plan",
        ["Phase", "Activities", "Duration / Week"],
        [
            ["1. Problem study", "Collect sample lures, study attack patterns, fix requirements", "Weeks 1-2"],
            ["2. System design", "Architecture, data-flow, use-case and database diagrams", "Weeks 3-4"],
            ["3. Analysis engines", "URL/domain heuristics, threat-feed client, psychology rules", "Weeks 5-7"],
            ["4. Decision and explanation", "Weighted fusion, feed minimum scores, Simple/Technical views", "Week 8"],
            ["5. Browser extension", "Navigation gate, cache, fail-closed handling, install guide", "Weeks 9-10"],
            ["6. Testing and documentation", "Automated tests, sample evaluation, deployment, report", "Weeks 11-12"],
        ],
        [4.6, 8.6, 2.7],
        center_cols=(2,),
        bold_first_col=True,
    )

    b.heading("4.8 Future Scope", level=2)
    b.add(
        "The present prototype relies on hand-written rules and public feeds. The next step is to assemble a "
        "labelled dataset of phishing and legitimate URLs and messages (for example from PhishTank, OpenPhish and "
        "public benchmark collections), train a supervised classifier on the features the heuristics already "
        "extract, and add its probability as one more weighted signal with precision, recall and F1 reported on a "
        "held-out set. Further extensions include text recognition for screenshots, a Firefox version of the "
        "extension, and migration to a server database if multi-user load grows."
    )

    # ------------------------------------------------------------------ 5
    b.heading("5. Tools / Technology Used")
    b.heading("5.1 Minimum Hardware Requirements", level=2)
    b.add(
        "The following configuration is sufficient to run the service, the development server, Chrome and the "
        "QR decoder together on a student laptop. No graphics card is used."
    )
    b.table(
        "Minimum hardware requirements",
        ["Component", "Minimum requirement"],
        [
            ["Processor", "64-bit dual-core, 2.0 GHz or faster (Intel Core i3 class or equivalent)"],
            ["Memory", "8 GB RAM (16 GB recommended when the browser, build tools and service run together)"],
            ["Storage", "10 GB free disk space for the code, Python environment, Node modules and the database"],
            ["Graphics", "Not required; integrated graphics is sufficient"],
            ["Network", "Internet connection for threat-feed lookups and deployment"],
            ["Peripherals", "Keyboard, display; a camera is not needed because QR images are uploaded as files"],
        ],
        [3.6, 12.3],
        bold_first_col=True,
    )

    b.heading("5.2 Minimum Software Requirements", level=2)
    b.add(
        "Versions are those declared in the project's dependency files; \">=\" denotes the minimum version "
        "accepted by the project and \"^\" the compatible range used by the front-end build."
    )
    b.table(
        "Software requirements and technology stack",
        ["Category", "Tool / Library", "Version", "Purpose"],
        [
            ["Language", "Python", "3.12.8", "Service and analysis engines"],
            ["Web framework / server", "FastAPI / Uvicorn [standard]", ">=0.115.0 / >=0.30.0", "REST JSON API and its server"],
            ["Validation", "Pydantic / pydantic-settings", ">=2.7.0 / >=2.2.0", "Request models and settings"],
            ["Authentication", "passlib[bcrypt] / bcrypt / python-jose", ">=1.7.4 / 4.0.1 / >=3.3.0", "Password hashing, login tokens"],
            ["HTTP clients", "requests / httpx", ">=2.31.0 / >=0.27.0", "Threat-feed and domain lookups"],
            ["Image processing", "opencv-python-headless / Pillow / NumPy", ">=4.8.0 / >=10.0.0 / >=1.26.0", "QR-code decoding from uploads"],
            ["Database", "SQLite (Python sqlite3 module)", "bundled with Python 3.12", "Accounts, scans, reports"],
            ["Other back-end", "python-dotenv / python-multipart / SQLAlchemy / pytest", ">=1.0.0 / >=0.0.9 / >=2.0.30 / >=8.0.0", "Configuration, uploads, ORM reserved for future work, tests"],
            ["Front-end runtime", "Node.js", "22.12.0", "Build tooling (deployment pin)"],
            ["UI and routing", "React, React DOM / react-router-dom / axios", "^19.2.8 / ^7.18.2 / ^1.19.0", "Web console, navigation, API calls"],
            ["Charts and build", "recharts, lucide-react / Vite, Tailwind CSS", "^3.10.1, ^1.31.0 / ^8.2.0, ^4.3.3", "Statistics charts, icons, bundling, styling"],
            ["Browser extension", "Chrome Extension Manifest V3", "extension 1.8.0", "Scan-before-open gate"],
            ["Platform", "Windows 10/11 or Linux; Chrome or other Chromium browser", "current", "Development, deployment, running the console"],
            ["Editor / optional", "Visual Studio Code or Cursor; Google Safe Browsing key; Tesseract OCR", "-", "Development; extra feed; text from screenshots"],
        ],
        [2.9, 5.0, 3.9, 4.1],
        size=10,
        bold_first_col=True,
    )

    # ------------------------------------------------------------------ 6
    b.heading("6. References")
    refs = [
        'R. M. Mohammad, F. Thabtah, and L. McCluskey, "Intelligent rule-based phishing websites classification," IET Information Security, vol. 8, no. 3, pp. 153-160, May 2014.',
        f'R. M. Mohammad, L. McCluskey, and F. Thabtah, "Phishing Websites," UCI Machine Learning Repository, 2015. [Online]. Available: https://archive.ics.uci.edu/dataset/327/phishing+websites. {ACCESSED}.',
        'J. Ma, L. K. Saul, S. Savage, and G. M. Voelker, "Beyond blacklists: Learning to detect malicious web sites from suspicious URLs," in Proc. 15th ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining (KDD), Paris, France, 2009, pp. 1245-1254.',
        'Y. Zhang, J. I. Hong, and L. F. Cranor, "CANTINA: A content-based approach to detecting phishing web sites," in Proc. 16th Int. Conf. World Wide Web (WWW), Banff, Canada, 2007, pp. 639-648.',
        'R. Dhamija, J. D. Tygar, and M. Hearst, "Why phishing works," in Proc. SIGCHI Conf. Human Factors in Computing Systems (CHI), Montreal, Canada, 2006, pp. 581-590.',
        'S. Sheng, B. Wardman, G. Warner, L. F. Cranor, J. Hong, and C. Zhang, "An empirical analysis of phishing blacklists," in Proc. 6th Conf. Email and Anti-Spam (CEAS), Mountain View, CA, USA, 2009.',
        'S. Garera, N. Provos, M. Chew, and A. D. Rubin, "A framework for detection and measurement of phishing attacks," in Proc. ACM Workshop on Recurring Malcode (WORM), Alexandria, VA, USA, 2007, pp. 1-8.',
        'A. K. Jain and B. B. Gupta, "A survey of phishing attack techniques, defence mechanisms and open research challenges," Enterprise Information Systems, vol. 16, no. 4, pp. 527-565, 2022.',
        f'Anti-Phishing Working Group, "Phishing Activity Trends Reports." [Online]. Available: https://apwg.org/trendsreports/. {ACCESSED}.',
        f'Cisco Talos, "PhishTank: Join the fight against phishing." [Online]. Available: https://phishtank.org/. {ACCESSED}.',
        f'abuse.ch, "URLhaus: Malware URL exchange." [Online]. Available: https://urlhaus.abuse.ch/. {ACCESSED}.',
        f'OpenPhish, "OpenPhish: Phishing intelligence." [Online]. Available: https://openphish.com/. {ACCESSED}.',
        f'S. Ramirez, "FastAPI." [Online]. Available: https://fastapi.tiangolo.com/. {ACCESSED}.',
        f'Google, "Chrome Extensions: Manifest V3." [Online]. Available: https://developer.chrome.com/docs/extensions/. {ACCESSED}.',
        f'SQLite Consortium, "SQLite Documentation." [Online]. Available: https://www.sqlite.org/docs.html. {ACCESSED}.',
    ]
    for i, ref in enumerate(refs, 1):
        p = b.add("", size=11, space_after=3, line="single")
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        r = p.add_run(f"[{i}]\t")
        _set_run_font(r, 11)
        r2 = p.add_run(ref)
        _set_run_font(r2, 11)


# --------------------------------------------------------------------------- #
# section / page setup
# --------------------------------------------------------------------------- #
def configure_section(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    _clear_runs(para)
    _add_page_field(para)
    _set_paragraph_format(para, size=11, align="right", space_after=0, line="single")
    for extra in footer.paragraphs[1:]:
        _delete_paragraph(extra)

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    for p in first_footer.paragraphs:
        _clear_runs(p)
    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    for p in first_header.paragraphs:
        _clear_runs(p)


def apply_document_defaults(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    rpr = normal._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def remove_trailing_empty(doc: Document):
    body = doc.element.body
    children = [c for c in body if c.tag != qn("w:sectPr")]
    for el in reversed(children):
        if el.tag == qn("w:p") and not "".join(el.itertext()).strip() and not el.findall(".//" + qn("w:drawing")):
            body.remove(el)
        else:
            break


# --------------------------------------------------------------------------- #
def ensure_figures() -> dict:
    names = [
        "fig1_system_architecture", "fig2_dfd_level0", "fig3_dfd_level1", "fig4_use_case",
        "fig5_er_diagram", "fig6_risk_decision_flow", "fig7_extension_gating_flow",
    ]
    missing = [n for n in names if not (FIGDIR / f"{n}.png").exists()]
    if missing:
        from render_figures import render_all
        render_all()
    fig = {n: FIGDIR / f"{n}.png" for n in names}
    for shot in ("fig8_scan_result", "fig9_history"):
        path = FIGDIR / f"{shot}.png"
        if not path.exists():
            path = _placeholder(shot)
        fig[shot] = path
    return fig


def _placeholder(name: str) -> Path:
    """Grey placeholder with instructions when a live screenshot is not available."""
    from PIL import ImageDraw, ImageFont
    path = FIGDIR / f"{name}_placeholder.png"
    im = Image.new("RGB", (2400, 1400), (234, 241, 248))
    d = ImageDraw.Draw(im)
    try:
        font = ImageFont.truetype("arial.ttf", 60)
    except OSError:
        font = ImageFont.load_default()
    msg = {
        "fig8_scan_result": "Screenshot placeholder: run the console, scan a URL and capture the result page.",
        "fig9_history": "Screenshot placeholder: open the History page while signed in and capture it.",
    }[name]
    d.rectangle([40, 40, 2360, 1360], outline=(11, 31, 58), width=8)
    d.text((120, 620), msg, fill=(11, 31, 58), font=font)
    im.save(path)
    return path


def build():
    template = _find_template()
    working = GENERATOR / "_working_template.docx"
    shutil.copy2(template, working)
    fig = ensure_figures()

    doc = Document(str(working))
    apply_document_defaults(doc)
    prepare_cover(doc)
    last_cover = strip_placeholders(doc)
    write_body(doc, last_cover, fig)
    remove_trailing_empty(doc)
    configure_section(doc)
    doc.save(str(OUTPUT))
    working.unlink(missing_ok=True)
    print(f"Wrote {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build()
